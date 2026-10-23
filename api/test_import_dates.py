"""Synthetic regression cases; no real input files or personal data."""
from datetime import datetime, time
from io import BytesIO

from django.contrib.auth.models import User
from django.test import SimpleTestCase, TestCase
from rest_framework.test import APIClient

from .models import Room, OperationLog, ScheduleVersion, ScheduleWriteLock
from .services.time_data import parse_time_entries
from .tests import make_xlsx_upload
from .views import normalize_import_cell


class ImportDateParsingTests(SimpleTestCase):
    def test_daily_date_range_with_comma(self):
        valid, invalid = parse_time_entries('2031-07-14至2031-07-16,09:00-18:00(每1天)')
        self.assertEqual(invalid, [])
        self.assertEqual(valid, [f'2031-07-{day} 09:00-18:00' for day in (14, 15, 16)])

    def test_chinese_date_and_full_width_clock(self):
        self.assertEqual(parse_time_entries('2031年7月14日 ９：００至１２：００'),
                         (['2031-07-14 09:00-12:00'], []))

    def test_hyphen_date_range(self):
        self.assertEqual(parse_time_entries('2031-07-14-2031-07-15 14:00-18:00'),
                         (['2031-07-14 14:00-18:00', '2031-07-15 14:00-18:00'], []))

    def test_invalid_or_unbounded_range_is_not_guessed(self):
        for value in ('7月14日 09:00-12:00', '2031-02-30 09:00-12:00',
                      '2031-07-16至2031-07-14 09:00-12:00',
                      '2031-01-01至2035-01-01 09:00-12:00',
                      '2031-07-14 18:00-09:00'):
            with self.subTest(value=value):
                valid, invalid = parse_time_entries(value)
                self.assertEqual(valid, [])
                self.assertTrue(invalid)

    def test_recurring_times_do_not_ignore_invalid_clocks_or_widen_explicit_hours(self):
        from .services.time_data import find_unrecognized_time_entries
        from .recurring_time import expand_recurring_entry
        from datetime import date
        self.assertTrue(find_unrecognized_time_entries('周一25:00-26:00'))
        self.assertTrue(find_unrecognized_time_entries('周一第7节'))
        self.assertEqual(expand_recurring_entry('周一下午15:00-16:00', date(2031, 7, 14), date(2031, 7, 14)),
                         ['2031-07-14 15:00-16:00'])

    def test_clock_cells_do_not_lose_their_time(self):
        self.assertEqual(normalize_import_cell(time(9, 30)), '09:30')
        self.assertEqual(normalize_import_cell(datetime(2031, 7, 14, 9, 30)), '2031-07-14 09:30')


class ImportPreviewTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(User.objects.create_user(username='import-admin', is_staff=True))

    def post(self, rows, **options):
        return self.client.post('/api/rooms/import_data/', {
            'file': make_xlsx_upload('rooms.xlsx', rows), **options,
        }, format='multipart')

    def test_preview_normalizes_periods_without_writing(self):
        version = ScheduleVersion.objects.create(defense_type='pre', is_current=True)
        lock_revision = ScheduleWriteLock.objects.get(pk=1).revision
        response = self.post([
            ['校区', '教室名称', '使用日期', '使用时间'],
            ['创新港', '测试楼-101', '2031-07-14至2031-07-15', '第1节-第8节'],
        ], preview='true')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(Room.objects.count(), 0)
        self.assertEqual(OperationLog.objects.count(), 0)
        version.refresh_from_db()
        self.assertEqual(version.revision, 0)
        self.assertEqual(ScheduleWriteLock.objects.get(pk=1).revision, lock_revision)
        self.assertTrue(response.data['canImport'])
        self.assertEqual(response.data['rows'][0]['row'], 2)
        self.assertIn('2031-07-15 14:00-18:00', response.data['rows'][0]['normalizedTime'])
        self.assertNotIn('08:00-18:00', response.data['rows'][0]['normalizedTime'])

    def test_card_range_and_explicit_campus(self):
        rows = [['借用记录'], ['状态：审核通过'],
                ['7-8101', '学院', '审核通过', '2031-07-14至2031-07-16,09:00-18:00(每1天)']]
        preview = self.post(rows, preview='true')
        self.assertFalse(preview.data['canImport'])
        self.assertIn('campus', preview.data['errors'][0]['errors'])
        result = self.post(rows, campus='创新港')
        self.assertEqual(result.status_code, 200, result.data)
        room = Room.objects.get()
        self.assertEqual(room.name, '7-8101')
        self.assertEqual(room.available_times.count('2031-07-'), 3)

    def test_original_row_numbers_survive_title_and_blank_rows(self):
        response = self.post([['教室资料'], [], ['校区', '教室名称', '可用时间'],
                              ['创新港', '测试楼-101', '日期待定']], preview='true')
        self.assertFalse(response.data['canImport'])
        self.assertEqual(response.data['errors'][0]['row'], 4)
        self.assertEqual(response.data['rows'][0]['row'], 4)

    def test_invalid_time_blocks_whole_import(self):
        response = self.post([['校区', '教室名称', '可用时间'],
                              ['创新港', '测试楼-101', '2031-07-14 09:00-12:00'],
                              ['创新港', '测试楼-102', '日期待定']])
        self.assertEqual(response.status_code, 400, response.data)
        self.assertEqual(Room.objects.count(), 0)

    def test_partial_booking_fields_do_not_become_unrestricted_availability(self):
        for day, clocks in [('2031-07-14', None), (None, '09:00-12:00')]:
            result = self.post([['校区', '教室名称', '使用日期', '使用时间'],
                                ['创新港', '测试楼-101', day, clocks]], preview='true')
            self.assertFalse(result.data['canImport'])

    def test_malformed_card_is_not_silently_skipped_after_a_valid_card(self):
        rows = [['借用记录'], ['7-8101', '2031-07-14 09:00-12:00'], ['7-8102', '日期待定']]
        response = self.post(rows, campus='创新港', preview='true')
        self.assertFalse(response.data['canImport'])
        self.assertEqual(response.data['errors'][0]['row'], 3)
        self.assertEqual(len(response.data['rows']), 2)

    def test_snake_case_headers_preserve_time_constraints(self):
        response = self.post([['campus', 'name', 'available_times'],
                              ['创新港', '测试楼-101', '2031-07-14 09:00-12:00']])
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(Room.objects.get().available_times, '2031-07-14 09:00-12:00')

    def test_missing_year_requires_explicit_choice(self):
        rows = [['校区', '教室名称', '使用日期', '使用时间'],
                ['创新港', '测试楼-101', '7月14日', '09:00-12:00']]
        self.assertFalse(self.post(rows, preview='true').data['canImport'])
        result = self.post(rows, importYear='2031')
        self.assertEqual(result.status_code, 200, result.data)
        self.assertEqual(Room.objects.get().available_times, '2031-07-14 09:00-12:00')

    def test_explicit_year_never_replaced(self):
        result = self.post([['校区', '教室名称', '使用日期', '使用时间'],
                            ['创新港', '测试楼-101', '2021-07-14', '09:00-12:00']], importYear='2031')
        self.assertEqual(result.status_code, 200, result.data)
        self.assertTrue(Room.objects.get().available_times.startswith('2021-07-14'))

    def test_merged_room_preview_keeps_both_sources(self):
        response = self.post([['校区', '教室名称', '使用日期', '使用时间'],
                              ['创新港', '测试楼-101', '2031-07-14', '09:00-12:00'],
                              ['创新港', '测试楼-101', '2031-07-15', '09:00-12:00']], preview='true')
        self.assertTrue(response.data['canImport'])
        self.assertEqual(response.data['rows'][0]['sourceRows'], [2, 3])
        self.assertEqual(response.data['validCount'], 1)

    def test_blank_unmerged_identity_is_not_inherited(self):
        response = self.post([['校区', '教室名称', '可用时间'],
                              ['创新港', '测试楼-101', '2031-07-14 09:00-12:00'],
                              [None, None, '2031-07-15 09:00-12:00']], preview='true')
        self.assertFalse(response.data['canImport'])
        self.assertEqual(response.data['errors'][0]['row'], 3)

    def test_actual_merged_cells_are_expanded_and_second_sheet_is_selectable(self):
        from openpyxl import Workbook
        from django.core.files.uploadedfile import SimpleUploadedFile
        workbook = Workbook()
        workbook.active.title = '空白说明'
        sheet = workbook.create_sheet('教室资料')
        sheet.append(['校区', '教室名称', '使用日期', '使用时间'])
        sheet.append(['创新港', '测试楼-101', '2031-07-14', '09:00-12:00'])
        sheet.append([None, '测试楼-102', None, '14:00-18:00'])
        sheet.merge_cells('A2:A3')
        sheet.merge_cells('C2:C3')
        buffer = BytesIO()
        workbook.save(buffer)
        def upload(**options):
            return self.client.post('/api/rooms/import_data/', {
                'file': SimpleUploadedFile('rooms.xlsx', buffer.getvalue()), **options,
            }, format='multipart')
        first = upload(preview='true')
        self.assertEqual(first.data['sheets'], ['空白说明', '教室资料'])
        self.assertFalse(first.data['canImport'])
        second = upload(preview='true', sheetName='教室资料')
        self.assertTrue(second.data['canImport'], second.data)
        self.assertEqual(second.data['rows'][1]['normalizedTime'], '2031-07-14 14:00-18:00')
        self.assertEqual(Room.objects.count(), 0)
        committed = upload(sheetName='教室资料')
        self.assertEqual(committed.status_code, 200, committed.data)
        self.assertEqual(Room.objects.count(), 2)

    def test_preview_requires_admin_permissions(self):
        self.client.force_authenticate(User.objects.create_user(username='import-reader'))
        result = self.post([['校区', '教室名称'], ['创新港', '测试楼-101']], preview='true')
        self.assertEqual(result.status_code, 403)
        self.assertEqual(Room.objects.count(), 0)
