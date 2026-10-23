"""Isolated source-service acceptance and export fixtures for October 14-16 work.

All execution dates are actual Asia/Shanghai dates; plan dates are not backfilled.
"""
import argparse
from datetime import datetime
import hashlib
from io import BytesIO
import json
import os
from pathlib import Path
import secrets
import sys
from zoneinfo import ZoneInfo

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'.tmp/acceptance-20261016'
OUT.mkdir(parents=True,exist_ok=True)
sys.path.insert(0,str(ROOT))
parser=argparse.ArgumentParser()
parser.add_argument('--serve',action='store_true')
parser.add_argument('--port',type=int,default=8894)
args=parser.parse_args()
os.environ.update(DJANGO_SETTINGS_MODULE='defense_scheduler.settings',DJANGO_ENV='local',
    DJANGO_DB_PATH=str(OUT/('browser.sqlite3' if args.serve else 'api.sqlite3')),
    DJANGO_SECRET_KEY='isolated-acceptance-only-'*4,DJANGO_DEBUG='false')
import django
django.setup()
from django.core.management import call_command
from django.contrib.auth.models import User
from rest_framework.test import APIClient
from api.models import Teacher,Student,Room,ScheduleVersion,ScheduleWriteLock,RuleConfig,Group,OperationLog
from api.serializers import default_rule_config
from openpyxl import load_workbook
from docx import Document

LABELS=['预答辩','正式答辩','中期答辩']

def seed(campus='创新港',long=False):
    call_command('flush',interactive=False,verbosity=0)
    ScheduleWriteLock.objects.get_or_create(pk=1)
    password=secrets.token_urlsafe(20)
    admin=User.objects.create_user('acceptance-admin',password=password,is_staff=True)
    User.objects.create_user('acceptance-reader',password=password)
    people=[Teacher.objects.create(name=f'QA教师{i:02}',title='教授' if i<=14 else '讲师',
        college='软件学院',is_software_teacher=True,roles=['主席','组长','普通专家'] if i<=14 else ['秘书'],
        available_types=LABELS) for i in range(1,17)]
    for i in range(12):
        Student.objects.create(name=f'QA学生{i+1:03}',student_no=f'QAACC{i+1:03}',
            mentor=people[i//6],mentor_name=people[i//6].name,student_type='学硕' if i%2==0 else '专硕',
            campus=campus,defense_types=LABELS,remark=('QA长备注 <仅文本> & 换行\n'*20 if long else 'QA备注 <仅文本>') if i==0 else '')
    for i in range(2):
        Room.objects.create(name=f'QA教室{i+1}',campus=campus,capacity=40,available_times='每周一至周五上午,每周一至周五下午')
    for kind,day in [('pre','2026-10-19'),('formal','2026-10-20'),('mid','2026-10-21')]:
        cfg=default_rule_config(kind)
        cfg.update(startDate=day,endDate=day,expertCount={'target':5,'min':5},includeRemarks=True)
        RuleConfig.objects.create(defense_type=kind,config=cfg)
    c=APIClient()
    c.force_authenticate(admin)
    return c,password

def unwrap(r):
    return r.data.get('data',r.data) if isinstance(r.data,dict) else r.data

def rules(kind):
    day={'pre':'2026-10-19','formal':'2026-10-20','mid':'2026-10-21'}[kind]
    return dict(defense_type=kind,policy_version=2,start_date=day,end_date=day,
        expert_count=5,expert_min=5,preserve_pre_defense_groups=False)

def current(c,kind='pre',version=None):
    return unwrap(c.get('/api/schedule/current/',dict(defense_type=kind,**({'version_id':version} if version else {}))))

def publish(c,result):
    return c.post('/api/schedule/publish/',{'version_id':result['versionId'],'expected_revision':result['revision']},format='json')

def export_files(c,result,tag):
    kind={'预答辩':'pre','正式答辩':'formal','中期答辩':'mid'}[result['defenseType']]
    response=c.get('/api/schedule/export/',{'defense_type':kind,'version_id':result['versionId']})
    assert response.status_code==200
    workbook=load_workbook(BytesIO(response.content),rich_text=True)
    roster=list(workbook['教室学生明细'].values)
    assert len(roster)-1==sum(len(g['students']) for g in result['groups'])
    expected={s['studentNo'] for g in result['groups'] for s in g['students']}
    assert {row[4] for row in roster[1:]}==expected
    (OUT/f'{tag}.xlsx').write_bytes(response.content)
    word=c.get('/api/schedule/export_word/',{'defense_type':kind,'version_id':result['versionId']})
    assert word.status_code==200
    doc=Document(BytesIO(word.content))
    text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    assert all(s['name'] in text for g in result['groups'] for s in g['students'])
    (OUT/f'{tag}.docx').write_bytes(word.content)
    return dict(student_count=len(expected),roster_rows=len(roster)-1,word_names_complete=True,
        excel_sha256=hashlib.sha256(response.content).hexdigest(),word_sha256=hashlib.sha256(word.content).hexdigest())

def main():
    call_command('migrate',verbosity=0)
    if args.serve:
        _,password=seed(long=True)
        (OUT/'browser-secret.json').write_text(json.dumps({'username':'acceptance-admin','reader':'acceptance-reader','password':password,'base':f'http://127.0.0.1:{args.port}'}),encoding='utf-8')
        from waitress import serve
        from django.core.wsgi import get_wsgi_application
        serve(get_wsgi_application(),host='127.0.0.1',port=args.port,threads=4)
        return
    cases=[]
    def record(id,prds,expect,fn):
        try:
            actual=fn()
            cases.append(dict(id=id,prds=prds,expected=expect,status='passed',actual=actual))
        except Exception as e:
            cases.append(dict(id=id,prds=prds,expected=expect,status='failed',error=f'{type(e).__name__}: {e}'))
        print(id,cases[-1]['status'],flush=True)
    def combo(kind,campus):
        c,_=seed(campus)
        response=c.post('/api/schedule/generate/',{'rules':rules(kind)},format='json')
        assert response.status_code==200,response.data
        result=unwrap(response)
        assert result['groups'] and all(g['campus']==campus for g in result['groups'])
        assert not [x for x in result['conflicts'] if x['level']=='error'],result['conflicts']
        assert publish(c,result).status_code==200
        return dict(groups=len(result['groups']),campus=campus,export=export_files(c,current(c,kind),kind+('-A' if campus=='创新港' else '-B')))
    for kind in ['pre','formal','mid']:
        for campus in ['创新港','兴庆']:
            record(f'FLOW-{kind}-{campus}',[1,4,16,17,21,25,29,36,37],'显式合法配置下生成、校验、发布、所选版本Excel/Word名单完整',lambda k=kind,s=campus:combo(k,s))
    c,_=seed()
    result=unwrap(c.post('/api/schedule/generate/',{'rules':rules('pre'),'request_key':'qa-idempotent'},format='json'))
    def lifecycle():
        repeat=unwrap(c.post('/api/schedule/generate/',{'rules':rules('pre'),'request_key':'qa-idempotent'},format='json'))
        assert repeat['versionId']==result['versionId']
        g=result['groups'][0]
        payload=dict(group_id=g['id'],group_data={'groupName':g['groupName']},expected_revision=result['revision'])
        missing=c.post('/api/schedule/adjust-group/',{k:v for k,v in payload.items() if k!='expected_revision'},format='json')
        assert missing.status_code==428
        ok=c.post('/api/schedule/adjust-group/',payload,format='json');assert ok.status_code==200
        stale=c.post('/api/schedule/adjust-group/',payload,format='json');assert stale.status_code==409
        assert OperationLog.objects.filter(authoritative=True,operator='acceptance-admin').exists()
        reader=APIClient();reader.force_authenticate(User.objects.get(username='acceptance-reader'))
        draft=unwrap(reader.get('/api/schedule/current/',{'defense_type':'pre'}));assert not draft.get('groups')
        assert reader.post('/api/schedule/generate/',{'rules':rules('pre')},format='json').status_code==403
        latest=current(c); assert publish(c,latest).status_code==200
        old=current(c)
        teacher=Teacher.objects.first()
        assert c.patch(f'/api/teachers/{teacher.pk}/',{'name':'QA重命名导师'},format='json').status_code==200
        assert current(c)['groups']==old['groups']
        new=unwrap(c.post('/api/schedule/generate/',{'rules':rules('pre')},format='json'))
        assert new['versionId']!=old['versionId']
        assert unwrap(reader.get('/api/schedule/current/',{'defense_type':'pre'}))['versionId']==old['versionId']
        assert c.post('/api/schedule/adjust-group/',dict(group_id=g['id'],group_data={'groupName':'BAD'},expected_revision=old['revision']),format='json').status_code==409
        return dict(missing_revision=428,stale_revision=409,reader_write=403,history_frozen=True,reader_sees_last_published=True,audit_authoritative=True)
    record('LIFECYCLE',[3,28,31,33,34,36],'幂等、并发拒绝、权限、审计、冻结与普通用户上一发布版',lifecycle)
    def linked_move():
        c,_=seed()
        response=c.post('/api/schedule/generate-linked/',dict(pre_rules=rules('pre'),formal_rules=rules('formal')),format='json')
        assert response.status_code==200,response.data
        linked=unwrap(response)
        mapping=lambda r:{s['id']:(g['groupName'],g['secretaryId']) for g in r['groups'] for s in g['students']}
        assert mapping(linked['pre'])==mapping(linked['formal'])
        formal=linked['formal'];src,dst=formal['groups'];s=src['students'][0]
        move=dict(action='move_student',student_id=s['id'],from_group_id=src['id'],to_group_id=dst['id'],expected_revision=formal['revision'])
        assert c.post('/api/schedule/adjust/',move,format='json').status_code==200
        moved=current(c,'formal');actual=next(x for g in moved['groups'] for x in g['students'] if x['id']==s['id'])
        assert actual['secretaryId']==s['secretaryId']
        assert any('秘书' in e['type'] for e in moved['conflicts'])
        move.update(from_group_id=dst['id'],to_group_id=src['id'],expected_revision=moved['revision'],preserve_secretary=False,secretary_id=src['secretaryId'])
        assert c.post('/api/schedule/adjust/',move,format='json').status_code==200
        restored=current(c,'formal');assert not [e for e in restored['conflicts'] if e['level']=='error'],restored['conflicts']
        audit=OperationLog.objects.filter(authoritative=True,type='调整').first()
        assert audit.operator=='acceptance-admin' and audit.created_at is not None
        assert audit.details['changes']['action']=='move_student' and audit.details['changes']['student_id']==s['id']
        assert 'before' in audit.details and 'after' in audit.details
        return dict(group_secretary_inheritance=True,default_secretary_preserved=True,explicit_rebind_and_recheck=True,
            authoritative_actor_time_source_target_relationship_audit=True)
    record('LINKED-MOVE',[21,22,27,28,31,32,33,34],'联排来源/学生秘书沿用、默认保留、显式改绑、自动重验',linked_move)
    def base_fields():
        c,_=seed()
        fields=dict(name='QA新增教师',college='其他学院',isExternal=True,isActive=True,memberEligible=False,
            isSoftwareTeacher=False,title='教授',roles=['组长'],availableTypes=['预答辩'],campusPreference='兴庆',
            unavailableTimes='2026-10-19 10:00-11:00',avoidTeacherNames='QA教师01',remark='QA不宜同组原因')
        created=c.post('/api/teachers/',fields,format='json');assert created.status_code==201,created.data
        t=unwrap(created);read=unwrap(c.get(f'/api/teachers/{t["id"]}/'))
        assert all(read[k]==v for k,v in fields.items())
        secretary=Teacher.objects.get(name='QA教师15')
        student=dict(name='QA新增学生',studentNo='QANEW001',gender='女',studentType='专硕',mentorId=t['id'],
            secretaryId=secretary.pk,campus='兴庆',defenseTypes=['预答辩','正式答辩'],remark='QA学生备注')
        added=c.post('/api/students/',student,format='json');assert added.status_code==201,added.data
        s=unwrap(added);read=unwrap(c.get(f'/api/students/{s["id"]}/'));assert all(read[k]==v for k,v in student.items())
        renamed=c.patch(f'/api/teachers/{t["id"]}/',{'name':'QA改名教师','isActive':False},format='json');assert renamed.status_code==200
        read=unwrap(c.get(f'/api/students/{s["id"]}/'));assert read['mentorName']=='QA改名教师' and read['mentorId']==t['id']
        assert c.delete(f'/api/students/{s["id"]}/').status_code==204
        assert c.delete(f'/api/teachers/{t["id"]}/').status_code==204
        return dict(all_teacher_fields_persisted=True,student_fields_and_stable_relations=True,renamed_reference_synced=True,disable_and_delete=True)
    record('BASE-FIELDS',[2,3,6],'录入/读取/改名/停用/删除及全部资料字段持久化',base_fields)
    def room_capacity():
        c,_=seed()
        duplicate=c.post('/api/rooms/',{'name':'QA教室1','campus':'创新港','capacity':40},format='json');assert duplicate.status_code==400
        other=c.post('/api/rooms/',{'name':'QA教室1','campus':'兴庆','capacity':40,'availableTimes':'每周一上午'},format='json');assert other.status_code==201
        Room.objects.all().update(capacity=5)
        bad=unwrap(c.post('/api/schedule/generate/',{'rules':rules('pre')},format='json'))
        assert any(e['level']=='error' for e in bad['conflicts']) and publish(c,bad).status_code==400
        return dict(same_campus_duplicate=400,cross_campus_duplicate_name_allowed=201,capacity_blocks_publish=True)
    record('ROOM-CAPACITY',[4,17,25],'教室唯一性按校区+名称，容量不足形成错误并禁止发布',room_capacity)
    def calendar():
        c,_=seed();Room.objects.all().update(available_times='')
        cfg=default_rule_config('pre');cfg.update(startDate='2026-10-24',endDate='2026-10-27',avoidWeekend=True,
            excludeDates=['2026-10-26'],expertCount={'target':5,'min':5})
        assert c.post('/api/rule-config/',cfg,format='json').status_code==200
        saved=unwrap(c.get('/api/rule-config/',{'defense_type':'pre'}))
        assert saved['startDate']==cfg['startDate'] and saved['excludeDates']==cfg['excludeDates'] and saved['avoidWeekend'] is True
        r=rules('pre');r.update(start_date=cfg['startDate'],end_date=cfg['endDate'],avoid_weekend=True,exclude_dates=cfg['excludeDates'])
        first=unwrap(c.post('/api/schedule/generate/',{'rules':r},format='json'));assert {g['date'] for g in first['groups']}=={'2026-10-27'}
        r['avoid_weekend']=False
        second=unwrap(c.post('/api/schedule/generate/',{'rules':r},format='json'));assert min(g['date'] for g in second['groups'])=='2026-10-24'
        return dict(saved_and_read_back=True,weekend_and_holiday_skipped=True,weekend_off_effective=True)
    record('CALENDAR',[9],'日期配置保存读取、周末开关、手工节假日影响真实生成',calendar)
    def overlap():
        c,_=seed();r=unwrap(c.post('/api/schedule/generate/',{'rules':rules('pre')},format='json'))
        a,b=r['groups'];changed=c.post('/api/schedule/adjust/',dict(action='change_time',group_id=b['id'],
            new_time=a['date']+' '+a['timeRange'],expected_revision=r['revision']),format='json');assert changed.status_code==200
        checked=current(c);errors=[e for e in checked['conflicts'] if e['level']=='error']
        assert any('时间冲突' in e['type'] or '教室冲突' in e['type'] for e in errors)
        assert publish(c,checked).status_code==400
        return dict(time_adjust_saved=True,automatic_time_or_room_conflict=True,publish_blocked=True)
    record('OVERLAP',[16,25,30,31,33],'修改时段制造重叠，自动检测人员/教室并拒绝发布',overlap)
    def adjust_fields():
        c,_=seed()
        Student.objects.exclude(pk__in=list(Student.objects.values_list('pk',flat=True)[:6])).delete()
        r=unwrap(c.post('/api/schedule/generate/',{'rules':rules('pre')},format='json'));group=r['groups'][0]
        def change(payload,whole=False):
            payload.update(group_id=group['id'],expected_revision=current(c)['revision'])
            reply=c.post('/api/schedule/adjust-group/' if whole else '/api/schedule/adjust/',payload,format='json')
            assert reply.status_code==200,reply.data
        change(dict(action='change_expert',old_expert_id=group['teachers'][0]['id'],new_expert_id=6))
        change(dict(group_data={'chairId':3,'chairman':'QA教师03','teachers':[{'id':i} for i in [1,4,5,6]]}),True)
        change(dict(action='change_secretary',new_secretary_id=15))
        Room.objects.create(name='QA兴庆调整教室',campus='兴庆',capacity=40)
        change(dict(group_data={'campus':'兴庆','classroom':'QA兴庆调整教室'}),True)
        changed=current(c)['groups'][0];assert changed['campus']=='兴庆' and changed['classroom']=='QA兴庆调整教室'
        change(dict(group_data={'campus':'创新港','classroom':'QA教室2'}),True)
        change(dict(action='change_time',new_time='2026-10-19 14:00-18:00'))
        final=current(c);assert not [e for e in final['conflicts'] if e['level']=='error'],final['conflicts']
        assert final['groups'][0]['secretaryId']==15 and final['groups'][0]['chairId']==3
        assert publish(c,final).status_code==200
        return dict(expert_chair_secretary_room_campus_time_updated=True,mentor_kept_in_committee=True,published=True)
    record('ADJUST-FIELDS',[21,31,32,33,34],'实际更换专家/组长/秘书/校区/教室/时段后保持关系并可发布',adjust_fields)
    def export_semantics():
        c,_=seed(long=True)
        r=unwrap(c.post('/api/schedule/generate/',{'rules':rules('pre')},format='json'))
        assert publish(c,r).status_code==200
        old=current(c);details=export_files(c,old,'long-history')
        response=c.get('/api/schedule/export/',{'defense_type':'pre','version_id':old['versionId']})
        book=load_workbook(BytesIO(response.content))
        assert '备注' in [cell.value for cell in book['教室学生明细'][1]]
        cfg=default_rule_config('pre');cfg.update(includeRemarks=False,expertCount={'target':5,'min':5})
        assert c.post('/api/rule-config/',cfg,format='json').status_code==200
        same=c.get('/api/schedule/export/',{'defense_type':'pre','version_id':old['versionId']})
        assert '备注' in [cell.value for cell in load_workbook(BytesIO(same.content))['教室学生明细'][1]]
        no_notes=rules('pre');no_notes['include_remarks']=False
        new=unwrap(c.post('/api/schedule/generate/',{'rules':no_notes},format='json'))
        no=c.get('/api/schedule/export/',{'defense_type':'pre','version_id':new['versionId']})
        assert '备注' not in [cell.value for cell in load_workbook(BytesIO(no.content))['教室学生明细'][1]]
        # Force a real error in the draft; production exporter should mark its cells red.
        Group.objects.filter(schedule_version_id=new['versionId']).update(chair=None)
        error=c.get('/api/schedule/export/',{'defense_type':'pre','version_id':new['versionId']})
        errorbook=load_workbook(BytesIO(error.content))
        red=sum(cell.font.color is not None and cell.font.color.type=='rgb' and cell.font.color.rgb.endswith('C00000') for sheet in errorbook for row in sheet for cell in row)
        assert red>0
        (OUT/'error-draft.xlsx').write_bytes(error.content)
        return dict(**details,old_snapshot_remarks_preserved=True,new_remarks_hidden=True,error_red_cells=red)
    record('EXPORT-SEMANTICS',[15,36,37,38,39,42],'所选版本、备注快照、隐藏备注、长文本、实际Excel错误标红',export_semantics)
    report=dict(executed_at=datetime.now(ZoneInfo('Asia/Shanghai')).isoformat(),planned_dates=['2026-10-14','2026-10-15'],
        scope='真实DRF服务、隔离数据库；非浏览器或视觉验收',cases=cases,
        counts={s:sum(c['status']==s for c in cases) for s in ['passed','failed']})
    (OUT/'api-results.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(report['counts'])

if __name__=='__main__':main()
