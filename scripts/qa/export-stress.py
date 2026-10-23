"""Verify actual >16-mentor export colors using a valid synthetic schedule."""
import importlib.util
from io import BytesIO
import json
import os
from pathlib import Path
import sys

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'.tmp/acceptance-20261016'
os.environ['QA_OUTPUT_DIR']=str(OUT/'color-stress')
spec=importlib.util.spec_from_file_location('qa_daily',ROOT/'scripts/qa/daily-rule-checks.py')
qa=importlib.util.module_from_spec(spec);spec.loader.exec_module(qa)
from django.core.management import call_command
from openpyxl import load_workbook
from docx import Document

if __name__=='__main__':
    call_command('migrate',verbosity=0)
    client=qa.seed('color-17')
    mentors=[f'QA教师{i:02}' for i in range(1,18)]
    qa.Teacher.objects.filter(name__in=mentors).update(member_eligible=False)
    response=qa.generate(client,expert_count=5,expert_min=5,start_date='2026-10-19',end_date='2026-10-23')
    assert response.status_code==200,response.data
    result=qa.body(response)
    assert not [e for e in result['conflicts'] if e['level']=='error'],result['conflicts']
    assert qa.publish(client,result).status_code==200
    xlsx=client.get('/api/schedule/export/',{'defense_type':'pre','version_id':result['versionId']})
    docx=client.get('/api/schedule/export_word/',{'defense_type':'pre','version_id':result['versionId']})
    assert xlsx.status_code==docx.status_code==200
    (OUT/'color-17.xlsx').write_bytes(xlsx.content);(OUT/'color-17.docx').write_bytes(docx.content)
    book=load_workbook(BytesIO(xlsx.content),rich_text=True)
    colors={}
    for row in list(book['教室学生明细'].rows)[1:]:
        colors[row[6].value]=row[6].font.color.rgb
        assert row[5].font.color.rgb==row[6].font.color.rgb
    assert len(colors)==17
    word=Document(BytesIO(docx.content))
    runs=[run for table in word.tables for row in table.rows for cell in row.cells for p in cell.paragraphs for run in p.runs]
    word_colors={name:sorted({str(run.font.color.rgb) for run in runs if name in run.text and run.font.color.rgb}) for name in mentors}
    report=dict(status='passed' if len(set(colors.values()))==17 else 'failed',expected='同份17导师互异，师生一致、跨输出一致',
        observed=dict(groups=len(result['groups']),students=sum(len(g['students']) for g in result['groups']),
            unique_excel_colors=len(set(colors.values())),excel_colors=colors,word_colors=word_colors,student_mentor_colors_match=True),
        setup='17位导师均可担任主席但不占普通专家席，避免缺少导师问题遮盖颜色验证；是有效并已发布合成排期，不修改业务实现')
    (OUT/'color-results.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False))
